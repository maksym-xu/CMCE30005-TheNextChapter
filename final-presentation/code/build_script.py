"""Build the rehearsal DOCX from script.json using the bundled python-docx runtime.

Run this file from any directory. Output defaults to ../deliverables.
Set FINAL_DOCX to an absolute or project-relative output path to preserve prior builds.
Render with the canonical documents/render_docx.py and bundled LibreOffice only.
"""
from pathlib import Path
import json
import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((ROOT / 'code/script.json').read_text())
requested_output = os.environ.get('FINAL_DOCX')
OUT = Path(requested_output).expanduser() if requested_output else ROOT / 'deliverables/TheNextChapter_Final_Speaking_Script.docx'
if not OUT.is_absolute():
    OUT = ROOT / OUT
OUT.parent.mkdir(parents=True, exist_ok=True)
assert len(DATA['slides']) == 10
assert sum(s['seconds'] for s in DATA['slides']) == DATA['target_seconds'] == 540
assert DATA['buffer_seconds'] == 60
assert DATA['target_seconds'] + DATA['buffer_seconds'] == 600
assert 880 <= DATA['spoken_word_count'] <= 950
assert DATA['spoken_word_count'] == sum(len(s['script'].split()) for s in DATA['slides'])
assert all(t['seconds'] == 135 and 200 <= t['words'] <= 255 for t in DATA['team_summary'])

DOC = Document()
section = DOC.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(.66)
section.bottom_margin = Inches(.66)
section.left_margin = Inches(.76)
section.right_margin = Inches(.76)
section.header_distance = Inches(.25)
section.footer_distance = Inches(.26)
for style in DOC.styles:
    if style.type == 1:
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
        rpr = style._element.get_or_add_rPr()
        color = rpr.find(qn('w:color'))
        if color is not None:
            for attr in ['themeColor', 'themeShade', 'themeTint']:
                color.attrib.pop(qn('w:' + attr), None)
        ppr = style._element.find(qn('w:pPr'))
        if ppr is not None:
            for border in list(ppr.findall(qn('w:pBdr'))):
                ppr.remove(border)
normal = DOC.styles['Normal']
normal.font.size = Pt(11.5)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.widow_control = True
DOC.styles['Title'].font.size = Pt(28)
DOC.styles['Title'].paragraph_format.space_after = Pt(18)
for name, size in [('Heading 1',19),('Heading 2',14),('Heading 3',12)]:
    st=DOC.styles[name]
    st.font.size=Pt(size)
    st.paragraph_format.space_before=Pt(10)
    st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.keep_with_next=True

header = section.header.paragraphs[0]
header.text = 'TheNextChapter Group 2   |   CMCE30005 final presentation'
header.runs[0].font.size = Pt(9)
footer=section.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Rehearsal script   ').font.size=Pt(9)
f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE');footer._p.append(f)

def para(text, size=None, bold=False, keep=False):
    p=DOC.add_paragraph()
    p.paragraph_format.keep_with_next=keep
    r=p.add_run(text);r.bold=bold;r.font.name='Arial'
    r.font.color.rgb=RGBColor(0,0,0)
    if size:r.font.size=Pt(size)
    if re.search('[\u4e00-\u9fff]',text):
        r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
        z=size or 10.5
        r.font.size=Pt(z)
        p.paragraph_format.line_spacing=Pt(z*1.4)
    return p

def heading(text,level=1):
    # DOCX headings use simple words and spaces; slide wording is retained in JSON.
    clean=re.sub(r'[^\w\s\u4e00-\u9fff]',' ',text)
    clean=re.sub(r'\s+',' ',clean).strip()
    return DOC.add_heading(clean,level)

def page_heading(text):
    # Keep the page break on the heading, avoiding a blank spillover paragraph.
    p=heading(text)
    p.paragraph_format.page_break_before=True
    return p

def table(rows,widths,size=10.5):
    t=DOC.add_table(rows=0,cols=len(widths));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for ci,v in enumerate(row):
            c=cells[ci];c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;c.text=str(v)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.05
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size=Pt(size);r.bold=ri==0
                    r.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
            pr=c._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'213547' if ri==0 else ('F2F4F6' if ri%2==0 else 'FFFFFF'));pr.append(sh)
            bd=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');bd.append(e)
            pr.append(bd)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
        if ri==0:
            trpr=t.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:tblHeader'))
        trpr=t.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
    return t

def time_label(sec):return f'{sec//60}:{sec%60:02d}'

DOC.add_paragraph('Melbourne Airbnb\nFinal presentation speaking script',style='Title')
para('英文演讲稿与中文排练提示  2026年10月1日版',16,True)
para('建议批准30天10套候选房源的集中搜索。优先调查City of Melbourne两至三卧公寓，将Yarra Ranges三卧住宅作为备选。每套房仍须通过许可和保守现金流检查。')
para('配套10页最终PPT。课堂只读英文正文；中文提示、问答和技术备查不计入演讲。')
heading('Run of show')
elapsed=0; rows=[['Presenter','Slides','Time','Words']]
for t in DATA['team_summary']:
    rows.append([t['speaker'],f"{t['slides'][0]} to {t['slides'][-1]}",f"{time_label(elapsed)} to {time_label(elapsed+t['seconds'])}",t['words']]);elapsed+=t['seconds']
table(rows,[2.15,1,1.9,1.2])
para(f"Total spoken text: {DATA['spoken_word_count']} words. Scripted time: {DATA['target_seconds']//60} minutes. Contingency: {DATA['buffer_seconds']} seconds. Q&A follows the presentation.",11)
para(DATA['word_count_method'],9)
heading('排练与交接')
para('四位组员各135秒。2:15由Qihang交给Maksym；4:30交给Eric；6:45交给Loc；9:00完成正文。留60秒处理换页、停顿或切换。正文已压缩，按自然语速预留指图和停顿。实际总时长由完整计时排练确定。')
para('每个人都要说得清Y、六个X、模型如何学权重产生历史排名分数，以及两种验证的区别。一次完整计时排练后，再练关键问答。3D示意图与淡入切换用于解释场景，不代表实际房源，也不是可自由旋转的三维模型。')
heading('必须准确区分的结果')
para('30条评论的模型验证与34条评论的跨年细分市场检验是两个不同设计。25.3%是分析样本观察结果，不是强制的25%。评论不是利润，当前设置可能不同于过去。')
para('在全部合资格房源中，相同历史搜索配额下，细分市场均值基准为28.1%，模型为43.7%，相差15.6个百分点。模型选中969套，其中423套达标，546套没有。30天和10套房是行动提案，不是模型求出的最优值。')

elapsed=0
for t in DATA['team_summary']:
    page_heading(t['speaker'])
    para(f"Slides {t['slides'][0]} to {t['slides'][-1]}   |   {time_label(elapsed)} to {time_label(elapsed+t['seconds'])}   |   {t['words']} spoken words",10,True)
    elapsed+=t['seconds']
    for s in [s for s in DATA['slides'] if s['speaker']==t['speaker']]:
        heading(f"Slide {s['slide']}  {s['title']}",2)
        para(f"Target: {s['seconds']} seconds",10,keep=True)
        para(s['script'])
        para('中文提示  '+s['zh_cue'],9.5)

for page in range(2):
    page_heading('Questions and answers' if page==0 else 'Questions and answers continued')
    para('以下内容用于问答，不在10分钟正文内朗读。先由指定主答人回答，其他组员只补充必要内容。',9.5)
    for n,q in enumerate(DATA['qa'][5*page:5*(page+1)],5*page+1):
        heading(f"{n}  {q['question']}",2)
        para('Lead answer: '+q['speaker'],9,True,keep=True)
        para(q['answer'],10.5)
        para('中文提示  '+q['zh_cue'],9)

page_heading('Technical details and source notes')
para('本页用于备查与回答追问，不放入课堂正文。模型复现结果与其他来源保存在同一交付目录的analysis子目录。',9)
nt=DATA['technical_notes']
heading('Model and test interpretation',2)
para(nt['outcome']+' '+nt['formula'],9.5)
para(nt['baseline'],9.5)
para(nt['ranking'],9.5)
para('Across all five host training folds, one-night acceptance is positively associated with the historical target, while higher relative price is negatively associated. These directions do not establish causal effects.',9)
para(nt['model_evolution'],9.5)
table([['Same input check','AUC','Brier','Screening rate'],['Logistic regression','0.7051','0.17180','43.65%'],['Fixed random forest','0.6838','0.18723','45.15%']],[2.3,1.05,1.15,1.75],9.5)
para(nt['matched_rf'],9)
para('Logistic confusion counts: selected 423 met / 546 missed; unselected 558 met / 2,346 missed. Precision is 43.7%; recall is 43.1%. Accuracy is not the screening objective.',9)
heading('Probability limits',2)
para('Individual probabilities remain uncertain. The highest score bin (above 60%) has 43 listings: 63.5% mean predicted versus 51.2% observed attainment. Future calibration is untested. Across 50 host splits, mean AUC is 0.708; the central 90% range of 0.649 to 0.778 is sensitivity, not a confidence interval.',9)
page_heading('Supplementary checks and sources')
para('本页区分推荐三组内的补充筛选检查与跨评论窗口比较。这些结果供追问时解释，主讲仍使用全部合资格样本的28.1%与43.7%。',9.5)
heading('Screening within the recommended groups',2)
para(nt['shortlist_scope_check'],10)
para(nt['shortlist_baseline_caution'],9.5)
para(nt['shortlist_uncertainty'],9.5)
heading('Comparing two review windows',2)
para('The separate temporal check has 2,657 surviving listings, 1,208 hosts, 10 segments and a 34-review cutoff. Both 365-day windows are reconstructed from dated reviews in the same supplied snapshot; they do not overlap. Earlier leaders retain a 12.3-point later-window lead, with a fixed-candidate host-bootstrap 95% range of 5.1 to 20.1 points. This is not a newly collected future snapshot and does not validate logistic forecasts. The slide 4 segment intervals overlap.',9.5)
heading('Sources and version',2)
para('Evidence and code: '+DATA['publication_url']+'. See EVIDENCE_MAP.md for slide-by-slide sources. Historical source commit: '+DATA['source_commit']+'.',9)
para('New results: analysis/outputs/presentation_metrics.json, model_comparison.csv, matched_rf_comparison.json, association_stability.json, calibration.csv and shortlisted_group_validation.json. Subgroup check: analysis/scripts/shortlisted_group_validation.py and independent verify_shortlisted_group_validation.R. Interim: analysis/reference/interim-model-selection. Temporal: analysis/reference/temporal-validation/validation_metrics.json. Models were rerun locally; temporal evidence was retained and checked, not rerun from raw reviews here.',9)
para('Course sources: Week 7 and Week 8 transcripts; Week 9 data storytelling; subject guide AI guidance. Individual sources appear in code/script.json and PowerPoint notes.',9)
heading('AI acknowledgement',2)
para(nt['ai_acknowledgement'],9)
DOC.core_properties.title=DATA['title']
DOC.core_properties.subject='CMCE30005 business presentation rehearsal and questions'
DOC.core_properties.author='TheNextChapter Group 2'
DOC.core_properties.keywords='Airbnb Melbourne presentation script'
DOC.save(OUT)
print(OUT)
print('Spoken words:',DATA['spoken_word_count'])
