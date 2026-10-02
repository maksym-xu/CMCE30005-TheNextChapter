"""Build the English rehearsal document and a plain-text speaking copy."""
from pathlib import Path
import json
import re
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

ROOT = Path(__file__).resolve().parents[1]
data = json.loads((ROOT / "code/script.json").read_text())
OUT = ROOT / "deliverables/TheNextChapter_English_Speaking_Script.docx"
assert len(data["slides"]) == 10 and len(data["qa"]) == 12
assert sum(s["seconds"] for s in data["slides"]) == 540
assert not re.search(r"[\u3400-\u9fff]", json.dumps(data, ensure_ascii=False))
assert 900 <= sum(len(s["script"].split()) for s in data["slides"]) <= 980

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin = section.bottom_margin = Inches(.66)
section.left_margin = section.right_margin = Inches(.76)
section.header_distance = Inches(.25)
section.footer_distance = Inches(.26)
for style in doc.styles:
    if style.type == 1:
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
        rpr = style._element.get_or_add_rPr()
        color = rpr.find(qn("w:color"))
        if color is not None:
            for name in ["themeColor", "themeShade", "themeTint"]:
                color.attrib.pop(qn("w:" + name), None)
        ppr = style._element.find(qn("w:pPr"))
        if ppr is not None:
            for border in list(ppr.findall(qn("w:pBdr"))):
                ppr.remove(border)
normal = doc.styles["Normal"]
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.widow_control = True
for name, size in [("Title", 26), ("Heading 1", 19), ("Heading 2", 14)]:
    style = doc.styles[name]
    style.font.size = Pt(size)
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.keep_with_next = True
header = section.header.paragraphs[0]
header.text = "TheNextChapter Group 2    CMCE30005 final presentation"
header.runs[0].font.size = Pt(9)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("English rehearsal script    ").font.size = Pt(9)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

def para(text, size=12, bold=False, keep=False):
    follows_table = len(doc.element.body) > 1 and doc.element.body[-2].tag == qn("w:tbl")
    p = doc.add_paragraph()
    if follows_table:
        p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.keep_with_next = keep
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(0, 0, 0)
    r.bold = bold
    return p

def heading(text, level=1, page=False):
    text = re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", text)).strip()
    p = doc.add_heading(text, level)
    p.paragraph_format.page_break_before = page
    return p

def table(rows, widths, size=10.5):
    t = doc.add_table(rows=0, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, value in enumerate(row):
            cell = cells[ci]
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = str(value)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.05
                for r in p.runs:
                    r.font.size = Pt(size)
                    r.bold = ri == 0
                    r.font.color.rgb = RGBColor(255, 255, 255) if ri == 0 else RGBColor(0, 0, 0)
            pr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "213547" if ri == 0 else ("F2F4F6" if ri % 2 == 0 else "FFFFFF"))
            pr.append(shade)
            borders = OxmlElement("w:tcBorders")
            for side in ["top", "left", "bottom", "right"]:
                border = OxmlElement("w:" + side)
                for key, val in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                    border.set(qn("w:" + key), val)
                borders.append(border)
            pr.append(borders)
            margins = OxmlElement("w:tcMar")
            for side in ["top", "left", "bottom", "right"]:
                el = OxmlElement("w:" + side)
                el.set(qn("w:w"), "90")
                el.set(qn("w:type"), "dxa")
                margins.append(el)
            pr.append(margins)
        rowpr = t.rows[-1]._tr.get_or_add_trPr()
        rowpr.append(OxmlElement("w:cantSplit"))
        if ri == 0:
            rowpr.append(OxmlElement("w:tblHeader"))
    return t

def link(label, url):
    p = doc.add_paragraph()
    element = OxmlElement("w:hyperlink")
    element.set(qn("r:id"), p.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "174E80"); props.append(color)
    size = OxmlElement("w:sz"); size.set(qn("w:val"), "20"); props.append(size)
    run.append(props)
    text = OxmlElement("w:t"); text.text = label; run.append(text)
    element.append(run); p._p.append(element)

def clock(seconds):
    return f"{seconds // 60}:{seconds % 60:02d}"

doc.add_paragraph("Melbourne Airbnb\nFinal presentation script", style="Title")
para("English edition   2 October 2026", 14, True)
para("We recommend a 30-day search for ten candidate homes, starting with City of Melbourne apartments and keeping Yarra Ranges houses as an alternative. Each candidate still needs permission and a conservative cash-flow assessment.")
para("This script follows the existing ten-slide presentation. Read only the spoken passages during the talk. Questions and evidence notes are preparation material.", 11)
heading("Presentation schedule")
rows = [["Presenter", "Slides", "Target time", "Words"]]
elapsed = 0
for speaker in data["team_summary"]:
    rows.append([speaker["speaker"], f"{speaker['slides'][0]} to {speaker['slides'][-1]}",
                 f"{clock(elapsed)} to {clock(elapsed + speaker['seconds'])}", speaker["words"]])
    elapsed += speaker["seconds"]
table(rows, [2.0, .9, 2.05, 1.8])
para(f"The spoken script has {data['spoken_word_count']} words. The target is nine minutes, leaving one minute for pauses and slide changes. Confirm the timing in a full rehearsal.", 11)
heading("Evidence to keep separate")
para("The slide 6 chart shows the original five-fold test: 43.7% versus 28.1%, with 423 of 969 selected homes reaching the target. The new 50-split comparison is a supplementary check: 43.42% versus 31.96% on average, with better screening in 48 splits. Its average gain is 11.46 percentage points.", 11)
para("The model uses a 30-review target. The separate comparison of two historical review windows uses 34. None of these results establishes future bookings or profit.", 11)

elapsed = 0
for speaker in data["team_summary"]:
    heading(speaker["speaker"], page=True)
    para(f"Slides {speaker['slides'][0]} to {speaker['slides'][-1]}    {clock(elapsed)} to {clock(elapsed + speaker['seconds'])}    {speaker['words']} spoken words", 10, True)
    elapsed += speaker["seconds"]
    for slide in [s for s in data["slides"] if s["speaker"] == speaker["speaker"]]:
        heading(f"Slide {slide['slide']}  {slide['title']}", 2)
        para(f"Target {slide['seconds']} seconds", 10, keep=True)
        para(slide["script"], 12.5)

for page in range(3):
    heading("Questions and answers" if page == 0 else "Questions and answers continued", page=True)
    para("Use these answers after the presentation. The named presenter leads; other members add only what is needed.", 10)
    for idx, question in enumerate(data["qa"][page * 4:(page + 1) * 4], page * 4 + 1):
        heading(f"{idx} {question['question']}", 2)
        para("Lead answer " + question["speaker"], 9, True, keep=True)
        para(question["answer"], 11)

heading("Evidence and sources", page=True)
para("The two model comparisons below use the same eligible population of 3,873 listings from 1,726 hosts. Their test designs and shortlist allocations differ.", 11)
heading("Original five fold comparison", 2)
para("The original test produced one held-out score per listing and pooled those scores for a 969-home shortlist. Logistic attainment was 43.65% against 28.07% for the segment-rate baseline, a gain of 15.58 percentage points. The revised preprocessing reproduced those results.", 11)
heading("Fifty repeated host splits", 2)
table([["Metric", "Logistic", "Baseline", "Better splits"],
       ["Shortlist target rate", "43.42%", "31.96%", "48 of 50"],
       ["AUC", "0.7084", "0.6035", "50 of 50"],
       ["Brier score", "0.17151", "0.18406", "44 of 50"]], [2.65, 1.3, 1.3, 1.5])
para("Each split used the same test hosts and the same shortlist quota for both methods. The mean paired target-rate gain was 11.46 percentage points. Its fifth to ninety-fifth percentiles were 2.37 to 19.72 points. These describe split sensitivity, not a confidence interval. Higher AUC and lower Brier are better.", 10.5)
para("The one-night indicator followed earlier outcome inspection. The split checks therefore evaluate a frozen exploratory specification. They do not repeat variable selection inside training or create an untouched external test. The random-forest, priority-group and historical-window results are earlier supplementary analyses.", 10.5)
para("The independent arithmetic check reconstructed scores from fitted coefficients and recalculated the metrics. It did not fit the model with a second engine or rebuild the original dated-review input.", 10.5)
base = "https://github.com/maksym-xu/CMCE30005-TheNextChapter/tree/final-presentation-release-2026-10-01/"
link("Repeated comparison code and aggregate results", base + "final-report/paired-analysis-2026-10-02")
link("Original presentation evidence and source map", base + "final-presentation")
link("English script source and deliverables", base + "final-presentation/english-script-2026-10-02")

doc.core_properties.title = "Melbourne Airbnb Final Presentation Script"
doc.core_properties.subject = "English presentation script and questions"
doc.core_properties.author = "TheNextChapter Group 2"
doc.core_properties.last_modified_by = "TheNextChapter Group 2"
doc.core_properties.comments = ""
doc.core_properties.keywords = "Melbourne Airbnb final presentation"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
# Remove unused non-English font labels inherited from the document template.
clean_copy = OUT.with_suffix(".clean.docx")
with ZipFile(OUT) as source, ZipFile(clean_copy, "w") as target:
    for info in source.infolist():
        payload = source.read(info.filename)
        if info.filename == "word/fontTable.xml":
            tree = etree.fromstring(payload)
            for element in list(tree):
                if any(re.search(r"[\u3400-\u9fff]", v) for v in element.attrib.values()):
                    tree.remove(element)
            payload = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
        elif info.filename == "word/theme/theme1.xml":
            tree = etree.fromstring(payload)
            for element in tree.iter():
                for name, value in list(element.attrib.items()):
                    if re.search(r"[\u3400-\u9fff]", value):
                        element.set(name, "Arial")
            payload = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
        target.writestr(info, payload)
clean_copy.replace(OUT)
spoken = ["Melbourne Airbnb Final Presentation Script", "English edition 2 October 2026", ""]
for s in data["slides"]:
    spoken.extend([f"Slide {s['slide']}  {s['title']}", f"{s['speaker']}  {s['seconds']} seconds", s["script"], ""])
(ROOT / "deliverables/TheNextChapter_English_Speaking_Script.txt").write_text("\n".join(spoken).rstrip() + "\n")
print(OUT.name)
print(f"Spoken words: {data['spoken_word_count']}")
