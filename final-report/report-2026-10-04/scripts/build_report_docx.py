"""Build Final_Report.docx from Final_Report.md using the report's Word template.

Run from the repository root:
    python final-report/report-2026-10-04/scripts/build_report_docx.py
Needs python-docx. The template (reference/report_template.docx) carries the page
size, margins, styles and page-number footer of the earlier Word report. The PDF is
exported from the .docx afterwards (LibreOffice: soffice --headless --convert-to pdf).
"""
from pathlib import Path
import re
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt, Inches, RGBColor

PKG = Path(__file__).resolve().parents[1]
MD = PKG / 'Final_Report.md'
TEMPLATE = PKG / 'reference/report_template.docx'
OUT = PKG / 'Final_Report.docx'
TEXT_WIDTH_TWIPS = 9866  # A4 width minus 2 x 1020 twip margins
FONT = 'Arial'


def set_run_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        fonts.set(qn(attr), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, text, url, size=None):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        fonts.set(qn(attr), FONT)
    rpr.append(fonts)
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0563C1')
    rpr.append(color)
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    rpr.append(underline)
    if size:
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), str(int(size * 2)))
        rpr.append(sz)
    run.append(rpr)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


TOKEN = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')


def add_inline(paragraph, text, size=None, bold=False):
    for piece in TOKEN.split(text):
        if not piece:
            continue
        if piece.startswith('**'):
            set_run_font(paragraph.add_run(piece[2:-2]), size, bold=True)
        elif piece.startswith('`'):
            set_run_font(paragraph.add_run(piece[1:-1]), size, bold=bold)
        elif piece.startswith('['):
            label, url = re.match(r'\[([^\]]+)\]\(([^)]+)\)', piece).groups()
            add_hyperlink(paragraph, label, url, size)
        elif piece.startswith('*'):
            set_run_font(paragraph.add_run(piece[1:-1]), size, bold=bold or None, italic=True)
        else:
            set_run_font(paragraph.add_run(piece), size, bold=bold or None)


def body_format(paragraph, after=6, before=None, line=1.08, keep_next=False):
    pf = paragraph.paragraph_format
    pf.space_after = Pt(after)
    if before is not None:
        pf.space_before = Pt(before)
    pf.line_spacing = line
    pf.widow_control = True
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if keep_next:
        pf.keep_with_next = True


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcpr.append(shd)


def borders(table):
    tblpr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:color'), 'D9D9D9')
        b.append(e)
    tblpr.append(b)
    mar = OxmlElement('w:tblCellMar')
    for edge in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:w'), '85')
        e.set(qn('w:type'), 'dxa')
        mar.append(e)
    tblpr.append(mar)
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'center')
    tblpr.append(jc)
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblpr.append(layout)


def plain(text):
    return re.sub(r'\*\*|\*|`', '', re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text))


def add_table(doc, rows, aligns):
    ncol = len(rows[0])
    # Each column first gets enough width for its longest word (so no word breaks),
    # then the remaining width is shared in proportion to content length.
    minimum, desired = [], []
    for c in range(ncol):
        cells = [plain(r[c]) for r in rows]
        longest_word = max((len(w) for t in cells for w in t.split()), default=4)
        numeric = [len(t) for t in cells[1:] if re.fullmatch(r'[-+0-9.,%–]+', t.strip())]
        need = max(longest_word * 125, max(numeric, default=0) * 115) + 200
        minimum.append(need)
        desired.append(max(len(t) for t in cells[1:]) if len(rows) > 1 else 10)
    spare = max(0, TEXT_WIDTH_TWIPS - sum(minimum))
    widths = [int(m + spare * d / sum(desired)) for m, d in zip(minimum, desired)]
    table = doc.add_table(rows=len(rows), cols=ncol)
    borders(table)
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(w))
    for i, row in enumerate(rows):
        tr = table.rows[i]
        trpr = tr._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for c, text in enumerate(row):
            cell = tr.cells[c]
            cell.width = widths[c] * 635  # twips to EMU
            tcw = cell._tc.get_or_add_tcPr().find(qn('w:tcW'))
            if tcw is not None:
                tcw.set(qn('w:type'), 'dxa')
                tcw.set(qn('w:w'), str(widths[c]))
            shade(cell, 'E9EEF3' if i == 0 else 'FFFFFF')
            p = cell.paragraphs[0]
            pf = p.paragraph_format
            pf.space_before = Pt(2)
            pf.space_after = Pt(2)
            pf.line_spacing = 1.0
            pf.alignment = WD_ALIGN_PARAGRAPH.RIGHT if (aligns[c] == 'right' and i > 0) else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, text.strip(), size=10.5, bold=(i == 0))
            if i < len(rows) - 1:
                pf.keep_with_next = True  # keep each table on one page
    return table


def parse_table(lines):
    rows = [[c.strip() for c in l.strip().strip('|').split('|')] for l in lines]
    spec = rows[1]
    aligns = ['right' if s.endswith(':') else 'left' for s in spec]
    return [rows[0]] + rows[2:], aligns


def remove_thumbnail(path):
    tmp = path.with_suffix('.tmp')
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename == 'docProps/thumbnail.jpeg':
                continue
            data = zin.read(item.filename)
            if item.filename == '_rels/.rels':
                data = re.sub(rb'<Relationship[^>]*thumbnail[^>]*/>', b'', data)
            if item.filename == '[Content_Types].xml':
                data = re.sub(rb'<Default Extension="jpeg"[^>]*/>', b'', data)
            zout.writestr(item, data)
    tmp.replace(path)


def build():
    doc = Document(str(TEMPLATE))
    body = doc.element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)
    lines = MD.read_text(encoding='utf-8').splitlines()
    in_refs = False
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith('|'):
            block = []
            while i < len(lines) and lines[i].startswith('|'):
                block.append(lines[i])
                i += 1
            rows, aligns = parse_table(block)
            add_table(doc, rows, aligns)
            spacer = doc.add_paragraph()
            body_format(spacer, after=4, line=1.0)
            continue
        if line.startswith('# '):
            p = doc.add_paragraph(style='Title')
            set_run_font(p.add_run(line[2:].strip()))
        elif line.startswith('## '):
            in_refs = line.startswith('## 8. References')
            p = doc.add_paragraph(style='Heading 1')
            set_run_font(p.add_run(line[3:].strip()))
            p.paragraph_format.keep_with_next = True
        elif line.startswith('### '):
            in_refs = False if line.startswith('### Author') or line.startswith('### AI') else in_refs
            p = doc.add_paragraph(style='Heading 2')
            set_run_font(p.add_run(line[4:].strip()))
            p.paragraph_format.keep_with_next = True
        elif line.startswith('!['):
            path = re.match(r'!\[[^\]]*\]\(([^)]+)\)', line).group(1)
            p = doc.add_paragraph()
            body_format(p, after=2, before=4, line=1.0, keep_next=True)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(PKG / path), width=Inches(6.6))
        elif re.match(r'\*\*Figure \d+\.\*\*', line):
            p = doc.add_paragraph(style='Caption')
            m = re.match(r'\*\*(Figure \d+\.)\*\*\s*(.*)', line)
            set_run_font(p.add_run(m.group(1)), bold=True)
            add_inline(p, ' ' + m.group(2))
        elif re.match(r'\*\*Table [A-Z]?\d+\.', line):
            p = doc.add_paragraph()
            body_format(p, after=4, before=4, keep_next=True)
            add_inline(p, line.strip('*'), bold=True)
        elif line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            body_format(p, after=4)
            add_inline(p, line[2:])
        elif re.match(r'\d+\. ', line):
            num, text = re.match(r'(\d+\.) (.*)', line).groups()
            p = doc.add_paragraph()
            body_format(p, after=4)
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            set_run_font(p.add_run(num + '\t'))
            add_inline(p, text)
        else:
            p = doc.add_paragraph()
            if in_refs:
                body_format(p, after=6, line=1.5)
                p.paragraph_format.left_indent = Inches(0.5)
                p.paragraph_format.first_line_indent = Inches(-0.5)
                p.paragraph_format.keep_together = True
            else:
                body_format(p)
            add_inline(p, line)
        i += 1
    props = doc.core_properties
    props.title = 'Melbourne Airbnb Market Screening'
    props.author = 'Eric Huang and Maksym Xu'
    props.last_modified_by = 'Group 2'
    props.subject = 'CMCE30005 Final Project Report'
    doc.save(str(OUT))
    remove_thumbnail(OUT)
    print('Wrote', OUT.relative_to(PKG.parents[1]))


if __name__ == '__main__':
    build()
