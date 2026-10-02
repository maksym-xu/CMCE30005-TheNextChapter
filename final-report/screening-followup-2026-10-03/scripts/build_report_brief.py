"""Create the English findings brief from checked aggregate tables."""
from pathlib import Path
import csv,re
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
def rows(name):
    with (ROOT/'outputs'/name).open() as f:return list(csv.DictReader(f))
sumrows=rows('screening_summary.csv'); pairs=rows('paired_screening_summary.csv')
def val(scope,frac,model,metric='precision'):
    return float(next(x['mean'] for x in sumrows if x['scope']==scope and float(x['fraction'])==frac and x['model']==model and x['metric']==metric))
def pair(scope,frac,name='full_vs_baseline'):
    return next(x for x in pairs if x['scope']==scope and float(x['fraction'])==frac and x['pair']==name and x['metric']=='precision_difference_pp')
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
for name in ['Normal','Title','Heading 1','Heading 2']:
    st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial')
    st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.line_spacing=1.08
normal=doc.styles['Normal'];normal.font.size=Pt(11)
for name,size in [('Title',24),('Heading 1',16),('Heading 2',13)]:
    st=doc.styles[name];st.font.size=Pt(size);st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(10)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run('TheNextChapter Group 2   ').font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
def p(t,size=11,bold=False):
    x=doc.add_paragraph();x.paragraph_format.widow_control=True
    r=x.add_run(t);r.font.size=Pt(size);r.bold=bold;return x
def heading(t):return doc.add_heading(t,1)
def table(values,widths):
    t=doc.add_table(rows=0,cols=len(widths));t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for ri,row in enumerate(values):
        cells=t.add_row().cells
        trpr=t.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,value in enumerate(row):
            c=cells[ci];c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;c.text=str(value)
            for para in c.paragraphs:
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_before=para.paragraph_format.space_after=Pt(4)
                para.paragraph_format.line_spacing=1.05
                for run in para.runs:run.font.size=Pt(10);run.bold=(ri==0);run.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
            pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'213547' if ri==0 else ('F1F4F3' if ri%2==0 else 'FFFFFF'));pr.append(sh)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side)
                for k,v in {'val':'single','sz':'4','color':'D9D9D9'}.items():e.set(qn('w:'+k),v)
                borders.append(e)
            pr.append(borders)
            mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

doc.add_paragraph('Additional screening evidence',style='Title')
p('Melbourne short-stay project   3 October 2026',11,True)
p('The full model added historical screening value after the three priority property groups had been chosen. Its usefulness depended partly on recorded price and minimum-stay information. A smaller shortlist was more concentrated, but missed more homes that reached the review target.')
heading('Value within the priority groups')
p('At a fixed 25% selection quota in each group, the full model achieved 49.6% target attainment against 32.4% for the baseline. The average gain was 17.2 percentage points, or about 17 additional target-reaching homes per 100 selected. It performed better in 46 of 50 host splits and worse in four. This gain cannot come from moving selection places between the groups.')
values=[['Evaluation scope','Group rate','Property only','Full model','Full gain\n(points)']]
for scope,label in [('all_eligible','All eligible listings'),('priority_pool','Three groups combined'),('priority_fixed_quota','Fixed quota in each group')]:
    z=pair(scope,.25)
    values.append([label,*[f'{100*val(scope,.25,m):.2f}%' for m in ['baseline','property_only','full']],f"+{float(z['mean_difference']):.2f}"])
table(values,[2.45,1.05,1.12,1.05,1.43])
p('All entries are equal-weight means across the same 50 overlapping host splits. The target is 30 or more reviews in the preceding 365 days. Quotas match between models within each row; fixed group quotas round up separately.',9.5)
p('The fifth-to-ninety-fifth percentile range for the fixed-group gain was -0.66 to 31.76 points. This describes variation across splits, not a confidence interval. The groups had already been selected during earlier exploration.',10.5)
heading('What the model needs to be useful')
p('Removing price and minimum-stay settings reduced fixed-group attainment to 38.6%. The full model added 11.0 points relative to this reduced model, with improvement in 42 splits and deterioration in eight. This is the joint predictive contribution of the two inputs after refitting the remaining weights, not the effect of changing either operating policy.')
p('The reduced model kept area, dwelling type, bedrooms and amenity count. Within each fixed group, only amenity count can distinguish homes. Results varied: in City of Melbourne three-bedroom apartments, the reduced model averaged 2.05 points below the baseline and was worse in 32 of 50 splits.')
p('Use the market groups to guide initial investigation. Use the full score as supplementary evidence for comparable established listings when recorded price, stay rules and other inputs are available. A proposed future price or stay policy is not a validated operating scenario. Permissions and conservative cash flow still determine whether a lease is acceptable.')

h=heading('Shortlist size and coverage');h.paragraph_format.page_break_before=True
p('The average gain remained positive at each tested quota. Choosing fewer listings improved concentration, while choosing more captured a larger share of all target-reaching homes. These checks do not identify an optimal search budget.')
values=[['Selected\nfrom each group','Full target\nattainment','Targets\ncaptured','Gain over\nbaseline','Better / tied /\nworse splits']]
for frac in [.1,.25,.5]:
    z=pair('priority_fixed_quota',frac)
    values.append([f'{frac:.0%}',f'{val("priority_fixed_quota",frac,"full"):.2%}',f'{val("priority_fixed_quota",frac,"full","recall"):.2%}',f'+{float(z["mean_difference"]):.2f} points',f'{z["better_splits"]} / {z["tied_splits"]} / {z["worse_splits"]}'])
table(values,[1.25,1.35,1.25,1.45,1.80])
p('Targets captured means the share of all target-reaching test listings included in the shortlist. At 10%, Yarra Ranges contributed only two to four selected listings per split, so individual group rates can be volatile.',9.5)
doc.add_picture(str(ROOT/'figures/quota_sensitivity.png'),width=Inches(7.0))
p('The evaluation retained the original 3,873 established listings with usable prices. Removing price as an input does not validate the reduced model for excluded unpriced homes or new rental candidates. Current settings may also differ from those used during the historical review period. These results do not establish future bookings, profit or causal effects.',10.5)
p('Sources: screening_summary.csv, paired_screening_summary.csv, paired_priority_group_summary.csv and priority_group_metrics.csv. Full definitions and reproducible code accompany this brief. The original five-fold presentation result remains separate.',9.5)

for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
for border in doc.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
for drawing in doc.element.xpath('.//wp:docPr'):
    drawing.set('descr', 'Mean historical screening gain at 10, 25 and 50 percent quotas, with fifth-to-ninety-fifth percentiles across 50 overlapping host splits. Full-model mean gains are 22.04, 17.19 and 10.67 percentage points.')

doc.core_properties.title='Additional Screening Evidence'
doc.core_properties.subject='Final report evidence on candidate screening'
doc.core_properties.author='TheNextChapter Group 2';doc.core_properties.last_modified_by='TheNextChapter Group 2';doc.core_properties.comments=''
out=ROOT/'deliverables/TheNextChapter_Screening_Findings.docx';doc.save(out)
clean=out.with_suffix('.clean.docx')
with ZipFile(out) as source,ZipFile(clean,'w') as target:
    for info in source.infolist():
        payload=source.read(info.filename)
        if info.filename in ['word/fontTable.xml','word/theme/theme1.xml']:
            tree=etree.fromstring(payload)
            for element in list(tree) if info.filename=='word/fontTable.xml' else []:
                if any(re.search(r'[\u3400-\u9fff]',v) for v in element.attrib.values()):tree.remove(element)
            for element in tree.iter():
                for name,value in list(element.attrib.items()):
                    if re.search(r'[\u3400-\u9fff]',value):element.set(name,'Arial')
            payload=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
        target.writestr(info,payload)
clean.replace(out)
print(out.name)
